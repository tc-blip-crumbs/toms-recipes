"""Turn an Eat Your Books bookshelf export into the site's cookbook index and a spreadsheet.

Usage: python3 tools/cookbooks.py eyb-bookshelf.json [cookbooks.xlsx]

The export is the JSON list collected from www.eatyourbooks.com/bookshelf, one object per recipe
with title, page, book, author, cat and ing. Writes src/cookbooks.json for the Cookbooks page and,
when a path is given, an .xlsx with one row per recipe and a Books tab.
"""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def main(src, xlsx=None):
    rows = json.loads(Path(src).read_text(encoding="utf-8"))
    seen, uniq = set(), []
    for r in rows:
        key = (r.get("title", ""), r.get("page", ""), r.get("book", ""))
        if key in seen or not r.get("title"):
            continue
        seen.add(key)
        uniq.append(r)
    books = []
    index = {}
    for r in uniq:
        k = (r["book"], r.get("author", ""))
        if k not in index:
            index[k] = len(books)
            books.append(list(k))
    order = sorted(range(len(books)), key=lambda i: books[i][0].lower())
    remap = {old: new for new, old in enumerate(order)}
    books = [books[i] for i in order]
    recipes = [[r["title"], r.get("page", ""), remap[index[(r["book"], r.get("author", ""))]], r.get("cat", ""), r.get("ing", "")] for r in uniq]
    recipes.sort(key=lambda x: (x[2], int(x[1]) if str(x[1]).isdigit() else 99999, x[0].lower()))
    out = ROOT / "src" / "cookbooks.json"
    out.write_text(json.dumps({"books": books, "recipes": recipes}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(recipes)} recipes in {len(books)} books -> {out}")

    if xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = "Recipes"
        ws.append(["Recipe", "Book", "Author", "Page", "Ingredients", "Categories"])
        for t, p, b, c, i in recipes:
            ws.append([t, books[b][0], books[b][1], int(p) if str(p).isdigit() else p, i.replace("; ", ", "), c])
        bs = wb.create_sheet("Books")
        bs.append(["Book", "Author", "Recipes"])
        for n, (title, author) in enumerate(books, start=2):
            bs.append([title, author, f'=COUNTIF(Recipes!B:B,A{n})'])
        head = PatternFill("solid", fgColor="1D1C19")
        for sheet, widths in ((ws, [48, 44, 22, 7, 70, 50]), (bs, [60, 26, 10])):
            for cell in sheet[1]:
                cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                cell.fill = head
                cell.alignment = Alignment(vertical="center")
            for col, w in zip("ABCDEF", widths):
                sheet.column_dimensions[col].width = w
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Arial")
        wb.save(xlsx)
        print(f"Spreadsheet -> {xlsx}")


if __name__ == "__main__":
    main(*sys.argv[1:])
